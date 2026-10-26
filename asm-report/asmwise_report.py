"""ASM-wise funnel report, split by lead source.

Logs in to River DMS as riverUserValidUser, exports the Enquiry List and Booking Master Tracker
Details for Yesterday, MTD and YTD into Output/Downloads/ASM, and builds
Output/ASM/ASMwise_AllSources_Summary_<yesterday>.xlsx with one block of rows per
Dealer Code + ASM and a Yesterday / MTD / YTD column group per source (see SOURCES):

    Dealer Code | ASM | Metric | Total Transactions Yesterday MTD YTD | Online ... | Activities ...

Bookings take the Lead Source of their enquiry (Enquiry Number = Lead No); a booking whose lead
isn't in the exports gets the usual Lead Source of its Primary Enquiry Source.

Usage (from the repo root):
    python asm-report/asmwise_report.py                   download the exports and build the summary
    python asm-report/asmwise_report.py --skip-download   rebuild from the last downloaded exports
    python asm-report/asmwise_report.py --keep-downloads  keep the exports under Output/Downloads/ASM

Environment: CREDENTIALS_FILE, ASM_OUTPUT_DIR, ASM_DOWNLOAD_DIR, BROWSER_CHANNEL, HEADLESS=1,
ASM_TODAY=YYYY-MM-DD (run as of that date).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import Page, TimeoutError as PWTimeout, sync_playwright

HERE = Path(__file__).resolve().parent  # asm-report
ROOT = HERE.parent
BASE_URL = "https://rivermobility.gaindms.com/"
CREDENTIALS_FILE = Path(os.environ.get("CREDENTIALS_FILE", ROOT / "resources" / "credentials.json"))
OUTPUT_DIR = Path(os.environ.get("ASM_OUTPUT_DIR", ROOT / "Output" / "ASM"))
# Its own folder: it is emptied after each run (clear_downloads).
DOWNLOAD_DIR = Path(os.environ.get("ASM_DOWNLOAD_DIR", ROOT / "Output" / "Downloads" / "ASM"))

# Column groups: (heading, Lead Source counted; None = every source). The Enquiry List form
# can't filter on Lead Source, so the exports are split here.
SOURCES: list[tuple[str, str | None]] = [
    ("Total Transactions", None), ("Online", "Online"), ("Walkin", "Walk-in"),
    ("Telephone", "Telephone"), ("Activities", "Activation"),
]
TOTAL = SOURCES[0][0]
SOURCE_LABEL = "All Sources"
SUMMARY_NAME = f"ASMwise_{SOURCE_LABEL.replace(' ', '')}_Summary"  # + _<dd-mm-yyyy>.xlsx
# The DMS renamed these Enquiry List columns in Oct 2026 (with its own spelling); accept both.
LEAD_NO = ("Lead No", "Oppurtunity No", "Opportunity No")
LEAD_SOURCE = ("Lead Source", "Primiry Source", "Primary Source")
SEC_SOURCE = ("Sec Lead Source", "Secoundry Source", "Secondary Source")
NO_TEST_RIDE = "Test Ride not taken"
RETAIL_STATUSES = {"delivered", "invoice", "invoiced"}
# Any other status (e.g. Cancelled) is counted as neither Booked-Active nor Retail.
BOOKED_ACTIVE_STATUSES = {"booked", "booked-active", "booked - active", "active", "booking", "open"}

METRICS = [
    "Enquiry", "Home Testride", "Store Testride", NO_TEST_RIDE,
    "Booked-Active", "Retail",
    "Enq to TR Conv%", "Enq to Booking Conv %", "Booking to Retails Conv %",
]
PERCENT_METRICS = {"Enq to TR Conv%", "Enq to Booking Conv %", "Booking to Retails Conv %"}
REPORTS = ("Enquiry List", "Booking Master Tracker")


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


@dataclass(frozen=True)
class Period:
    label: str
    start: date
    end: date


def periods(today: date) -> list[Period]:
    yesterday = today - timedelta(days=1)
    fy_start = date(yesterday.year if yesterday.month >= 4 else yesterday.year - 1, 4, 1)  # FY starts 1 April
    return [
        Period("Yesterday", yesterday, yesterday),
        Period("MTD", yesterday.replace(day=1), yesterday),
        Period("YTD", fy_start, yesterday),
    ]


def column(group: str, period_label: str) -> str:
    return f"{group} {period_label}"


def columns(plist: list[Period]) -> list[str]:
    """The value columns, source group by source group."""
    return [column(g, p.label) for g, _ in SOURCES for p in plist]


def export_name(name: str, period: Period) -> str:
    return f"{name.replace(' ', '_')}_{period.label}_{period.start:%Y%m%d}_{period.end:%Y%m%d}"


class Portal:
    def __init__(self, page: Page):
        self.page = page
        self.content = page.locator("#TableContent")

    def login(self, creds: dict) -> None:
        log("Logging in as River User...")
        p = self.page
        p.goto(BASE_URL, wait_until="domcontentloaded")
        p.get_by_role("link", name="River User").click()
        p.fill("#UserName", creds["username"])
        p.press("#UserName", "Tab")
        p.fill("#Password", creds["password"])
        p.click("#CmdLogin")
        try:
            p.locator("#ctl00_cpMain_DashboardCaption").wait_for(timeout=60000)
        except PWTimeout:
            raise SystemExit("Login failed: dashboard did not load (check riverUserValidUser in credentials.json).")
        log("Logged in.")

    def _menu(self, item_name: str) -> None:
        # MIS toggles open/closed, so only click it when the item isn't showing.
        item = self.page.get_by_role("link", name=item_name)
        if not item.is_visible():
            self.page.get_by_role("link", name=" MIS ").click()
            try:
                item.wait_for(state="visible", timeout=10000)
            except PWTimeout:
                self.page.get_by_role("link", name=" MIS ").click()
        item.click()

    def _fresh_form(self) -> None:
        # Reopen from the dashboard to reset the postback form.
        self.page.locator('a[href*="UserDashboard.aspx"]').first.click()
        self.page.wait_for_load_state("domcontentloaded")  # "load" can hang on slow dashboard assets
        self._menu(" Enquiry Reports")  # collapses any open submenu
        self._menu(" MFR Vehicles Reports")

    def open_enquiry_list(self) -> None:
        log("MIS > MFR Vehicles Reports > Pre-sales > Enquiry List")
        self._fresh_form()
        self.content.get_by_text("Pre-sales").click()
        self.content.get_by_text("Enquiry ListThe list of all").click()
        self.page.locator("#ctl00_cpMain_FromDate").wait_for()

    def open_booking_tracker(self) -> None:
        log("MIS > MFR Vehicles Reports > Sales > Booking Master Tracker Details")
        self._fresh_form()
        self.content.get_by_text("Sales", exact=True).click()
        self.content.get_by_text("Booking Master Tracker Details").first.click()
        self.page.locator("#ctl00_cpMain_FromDate").wait_for()
        all_status = self.page.locator("#ctl00_cpMain_BookingStatus_0")
        if all_status.is_visible():
            all_status.check()

    def set_date(self, input_id: str, d: date) -> None:
        # The calendar opens on the current month. The grid also shows days from the
        # neighbouring months, so the day is matched on its title ("Wednesday, April 1, 2026").
        p = self.page
        p.locator(f"#{input_id}").click()
        month, target = d.strftime("%B"), (d.year, d.month)
        title = p.locator(".ajax__calendar_title:visible")
        text = ""
        for _ in range(36):
            text = re.sub(r"\s+", " ", title.inner_text().replace(",", "")).strip()
            shown = datetime.strptime(text, "%B %Y")
            if (shown.year, shown.month) == target:
                break
            p.locator(".ajax__calendar_prev:visible" if (shown.year, shown.month) > target
                      else ".ajax__calendar_next:visible").click()
            # The title updates after the slide animation; reading it too early made the
            # loop click again on the stale month and overshoot back and forth.
            for _ in range(50):
                p.wait_for_timeout(100)
                if re.sub(r"\s+", " ", title.inner_text().replace(",", "")).strip() != text:
                    break
        else:
            raise RuntimeError(f"Calendar never reached {month} {d.year} (stuck on {text})")
        p.locator(f'.ajax__calendar_day:visible[title*="{month}"]').filter(
            has_text=re.compile(rf"^{d.day}$")).click()

    def export(self, period: Period, name: str) -> Path:
        self.set_date("ctl00_cpMain_FromDate", period.start)
        self.set_date("ctl00_cpMain_ToDate", period.end)
        log(f"{name} {period.label}: {period.start:%d-%m-%Y} to {period.end:%d-%m-%Y} - View Report")

        with self.page.expect_popup(timeout=300000) as popup_info:
            self.page.get_by_role("link", name=" View Report").click()
        viewer = popup_info.value
        downloads = []
        viewer.on("download", lambda d: downloads.append(d))
        viewer.on("dialog", lambda d: d.accept())
        # Headless, Export to Excel can hand the file over in another window; catch it there too.
        context = viewer.context

        def watch(pg: Page) -> None:
            pg.on("download", lambda d: downloads.append(d))
            pg.on("dialog", lambda d: d.accept())

        context.on("page", watch)
        try:
            download = self._wait_for_export(viewer, downloads)
            target = DOWNLOAD_DIR / f"{export_name(name, period)}{Path(download.suggested_filename).suffix or '.xlsx'}"
            target.parent.mkdir(parents=True, exist_ok=True)
            download.save_as(target)
            log(f"Saved {target.name}")
            return target
        finally:
            context.remove_listener("page", watch)
            for pg in context.pages:
                if pg != self.page:
                    pg.close()

    def _wait_for_export(self, viewer: Page, downloads: list):
        # Over 1000 rows the viewer shows a "record display limit" message and downloads the
        # Excel itself; otherwise it shows the grid with an "Export to Excel" button.
        # A big report (YTD) keeps the viewer busy loading, and a visibility check can then hang
        # until the 30s default timeout; treat that as "not ready yet" instead of failing.
        def showing(locator) -> bool:
            try:
                return locator.is_visible(timeout=5000)
            except PWTimeout:
                return False

        def wait_download(seconds: int) -> bool:
            end = time.time() + seconds
            while time.time() < end:
                if downloads:
                    return True
                viewer.wait_for_timeout(1000)
            return False

        # The download is caught by the listeners in export(), from whichever window it comes.
        # A click on Export to Excel while the grid is still settling can be ignored, so the
        # button is clicked again (up to 3 times) when no file arrives.
        deadline = time.time() + 420
        clicks = 0
        while time.time() < deadline:
            if downloads:
                return downloads[0]
            if showing(viewer.get_by_text(re.compile("record display limit", re.I))):
                if wait_download(15):
                    return downloads[0]
                viewer.get_by_role("link", name="click here").click(no_wait_after=True)
                if wait_download(180):
                    return downloads[0]
                raise RuntimeError("Clicked 'click here' but no Excel file arrived within 3 minutes.")
            export_btn = viewer.get_by_role("button", name="Export to Excel")
            if clicks < 3 and showing(export_btn):
                clicks += 1
                export_btn.click(no_wait_after=True)
                if wait_download(90):
                    return downloads[0]
                log(f"No file after Export to Excel (try {clicks} of 3)")
                continue
            viewer.wait_for_timeout(1000)
        raise RuntimeError("Report viewer produced no Excel download within 7 minutes.")


def download_reports(today: date) -> dict[tuple[str, str], Path]:
    creds = json.loads(CREDENTIALS_FILE.read_text(encoding="utf-8"))["riverUserValidUser"]
    files: dict[tuple[str, str], Path] = {}
    with sync_playwright() as pw:
        # Same flags as playwright.config.ts: Chrome blocks the .xls downloads and the
        # Export to Excel popup without them, and the viewer renders blank with a fixed viewport.
        browser = pw.chromium.launch(
            channel=os.environ.get("BROWSER_CHANNEL", "chrome"),
            headless=os.environ.get("HEADLESS") == "1",
            args=["--start-maximized", "--safebrowsing-disable-download-protection",
                  "--disable-client-side-phishing-detection", "--disable-popup-blocking"],
        )
        context = browser.new_context(no_viewport=True, accept_downloads=True)
        context.set_default_timeout(30000)
        portal = Portal(context.new_page())
        portal.login(creds)
        openers = {"Enquiry List": portal.open_enquiry_list, "Booking Master Tracker": portal.open_booking_tracker}
        for name in REPORTS:
            for period in periods(today):
                openers[name]()
                files[(name, period.label)] = portal.export(period, name)
        browser.close()
    return files


def latest_files(today: date) -> dict[tuple[str, str], Path]:
    files = {}
    for name in REPORTS:
        for period in periods(today):
            pattern = f"{export_name(name, period)}.*"
            found = sorted(DOWNLOAD_DIR.glob(pattern), key=lambda f: f.stat().st_mtime)
            if not found:
                raise SystemExit(f"--skip-download: no {pattern} in {DOWNLOAD_DIR}. Run without --skip-download first.")
            files[(name, period.label)] = found[-1]
    return files


def clean(name) -> str:
    return re.sub(r"\s+", " ", str(name).replace("\xa0", " ")).strip()


def key(name) -> str:
    return re.sub(r"[\s_]+", "", str(name).replace("\xa0", " ")).lower()


def _excel_engine() -> str | None:
    try:
        import python_calamine  # noqa: F401  much faster on the large YTD export
        return "calamine"
    except ImportError:
        return None


def read_export(path: Path) -> pd.DataFrame:
    """Reads a DMS export (.xlsx, or the HTML table saved as .xls for small reports),
    starting at the row that holds the Dealer Code header."""
    if path.read_bytes()[:200].lstrip().lower().startswith((b"<html", b"<!doctype", b"<table")):
        table = max(pd.read_html(path, header=None, keep_default_na=False), key=len).astype(str)
        # read_html turns the header row into column names; add it back as a row.
        raw = pd.concat([pd.DataFrame([list(map(str, table.columns))]), pd.DataFrame(table.values)],
                        ignore_index=True)
    else:
        raw = pd.read_excel(path, header=None, dtype=str, engine=_excel_engine())
    raw = raw.where(raw.notna() & (raw != "nan") & (raw != ""), None)
    header_row = next((i for i, row in raw.iterrows() if any(key(c) == "dealercode" for c in row if c)), None)
    if header_row is None:
        raise ValueError(f"No 'Dealer Code' header row in {path.name}")
    df = raw.iloc[header_row + 1:].copy()
    df.columns = [clean(c) if c else f"col{i}" for i, c in enumerate(raw.iloc[header_row])]
    df = df.dropna(how="all")
    return df.apply(lambda s: s.map(lambda v: clean(v) if v is not None else ""))


def col(df: pd.DataFrame, *names: str) -> str:
    by_key = {key(c): c for c in df.columns}
    for n in names:
        if key(n) in by_key:
            return by_key[key(n)]
    raise KeyError(f"None of {names} in columns: {list(df.columns)}")


def norm_dealer(v: str) -> str:
    return re.sub(r"\.0$", "", v.strip())


def enquiry_counts(df: pd.DataFrame) -> pd.DataFrame:
    dealer, asm, lead = col(df, "Dealer Code"), col(df, "ASM Name"), col(df, *LEAD_NO)
    tr = col(df, "Testride required at")
    df = df.assign(**{dealer: df[dealer].map(norm_dealer),
                      tr: df[tr].replace("", NO_TEST_RIDE)})
    # A lead can have several rows; count it once as Home, else Store, else not taken.
    rank = {"Home Testride": 0, "Store Testride": 1}
    df = (df.assign(_rank=df[tr].map(lambda v: rank.get(v, 2)))
            .sort_values("_rank").drop_duplicates([dealer, asm, lead]))
    grouped = (df.groupby([dealer, asm, tr])[lead].nunique()
                 .unstack(tr, fill_value=0)
                 .reindex(columns=["Home Testride", "Store Testride", NO_TEST_RIDE], fill_value=0))
    grouped["Enquiry"] = df.groupby([dealer, asm])[lead].nunique()
    grouped.index.names = ["Dealer Code", "ASM"]
    return grouped


def lead_sources(enquiries: list[pd.DataFrame]) -> tuple[dict[str, str], dict[str, str]]:
    """Lead No -> Lead Source, and Sec Lead Source (casefolded) -> its most common Lead Source."""
    df = pd.concat([e[[col(e, *LEAD_NO), col(e, *LEAD_SOURCE), col(e, *SEC_SOURCE)]]
                    .set_axis(["lead", "src", "sec"], axis=1) for e in enquiries])
    df = df[df["src"] != ""]
    by_lead = df.drop_duplicates("lead").set_index("lead")["src"].to_dict()
    sec = df[df["sec"] != ""]
    by_sec = sec.groupby(sec["sec"].str.casefold())["src"].agg(lambda s: s.mode().iat[0]).to_dict()
    return by_lead, by_sec


def booking_sources(df: pd.DataFrame, by_lead: dict[str, str], by_sec: dict[str, str]) -> pd.Series:
    # The tracker's Primary Enquiry Source is the enquiry's Sec Lead Source (Meta, Google, ...),
    # so a booking takes its lead's Lead Source, else the usual one for that secondary source.
    lead = df[col(df, "Enquiry Number", "Enquiry No")].map(by_lead)
    sec = df[col(df, "Primary Enquiry Source")].str.casefold().map(by_sec)
    return lead.fillna(sec).fillna("")


def booking_counts(df: pd.DataFrame) -> tuple[pd.DataFrame, set[str]]:
    dealer, asm, status = col(df, "Dealer Code"), col(df, "ASM Name"), col(df, "Status")
    order = col(df, "Sales Order", "Sales order")

    def bucket(s: str) -> str:
        s = s.strip().casefold()
        if s in RETAIL_STATUSES:
            return "Retail"
        if s in BOOKED_ACTIVE_STATUSES:
            return "Booked-Active"
        return "Other"

    df = df.assign(**{dealer: df[dealer].map(norm_dealer), "_bucket": df[status].map(bucket)})
    other = set(df.loc[df["_bucket"] == "Other", status].unique())
    grouped = (df[df["_bucket"] != "Other"].groupby([dealer, asm, "_bucket"])[order].nunique()
                 .unstack("_bucket", fill_value=0)
                 .reindex(columns=["Booked-Active", "Retail"], fill_value=0))
    grouped.index.names = ["Dealer Code", "ASM"]
    return grouped, other


def pct(num, den):
    return (num / den.where(den != 0) * 100).astype(float).round(1)


def add_conversions(t: pd.DataFrame) -> None:
    t["Enq to TR Conv%"] = pct(t["Home Testride"] + t["Store Testride"], t["Enquiry"])
    t["Enq to Booking Conv %"] = pct(t["Booked-Active"], t["Enquiry"])
    t["Booking to Retails Conv %"] = pct(t["Retail"], t["Booked-Active"] + t["Retail"])


def metrics(enq: pd.DataFrame, bkg: pd.DataFrame) -> pd.DataFrame:
    m = enq.join(bkg, how="outer").fillna(0).astype(int)
    add_conversions(m)
    return m[METRICS]


def build_summary(files: dict[tuple[str, str], Path], today: date) -> pd.DataFrame:
    plist = periods(today)
    per_column = {}
    enquiries = {}
    for period in plist:
        log(f"Reading {period.label} exports...")
        enquiries[period.label] = read_export(files[("Enquiry List", period.label)])
    by_lead, by_sec = lead_sources(list(enquiries.values()))

    for period in plist:
        enq_df = enquiries[period.label]
        bkg_df = read_export(files[("Booking Master Tracker", period.label)])
        enq_src = enq_df[col(enq_df, *LEAD_SOURCE)].str.casefold()
        bkg_src = booking_sources(bkg_df, by_lead, by_sec).str.casefold()
        log(f"{period.label}: {len(enq_df)} enquiry rows, {len(bkg_df)} booking rows "
            f"({int((bkg_src == '').sum())} bookings with no lead source - Total only)")
        for group, source in SOURCES:
            e = enq_df if source is None else enq_df[enq_src == source.casefold()]
            b = bkg_df if source is None else bkg_df[bkg_src == source.casefold()]
            enq, (bkg, other) = enquiry_counts(e), booking_counts(b)
            if source is None and other:
                log(f"[WARN] {period.label}: booking statuses counted as neither Booked-Active nor Retail: {sorted(other)}")
            log(f"  {group}: {int(enq['Enquiry'].sum())} leads, "
                f"{int(bkg['Booked-Active'].sum())} Booked-Active, {int(bkg['Retail'].sum())} Retail")
            per_column[column(group, period.label)] = metrics(enq, bkg)

    # One row per Dealer Code / ASM / metric; every pair gets all metric rows so the blocks line up.
    pairs = sorted(set().union(*(m.index for m in per_column.values())))
    full = pd.MultiIndex.from_tuples([(d, a, k) for d, a in pairs for k in METRICS],
                                     names=["Dealer Code", "ASM", "Metric"])
    long = pd.concat({label: m.stack(future_stack=True) for label, m in per_column.items()}, axis=1)
    long = long.reindex(full).reset_index()
    long["_order"] = long["Metric"].map(METRICS.index)
    long = long.sort_values(["Dealer Code", "ASM", "_order"]).drop(columns="_order")
    count_cols = columns(plist)
    is_count = ~long["Metric"].isin(PERCENT_METRICS)
    long.loc[is_count, count_cols] = long.loc[is_count, count_cols].fillna(0)
    return long[["Dealer Code", "ASM", "Metric", *count_cols]].reset_index(drop=True)


def overall(summary: pd.DataFrame, labels: list[str]) -> pd.DataFrame:
    """All-dealer totals; conversions are worked out from the totals."""
    t = summary[~summary["Metric"].isin(PERCENT_METRICS)].groupby("Metric")[labels].sum().T
    add_conversions(t)
    out = t[METRICS].T.reset_index(names="Metric")
    out.insert(0, "ASM", "")
    out.insert(0, "Dealer Code", "All Dealers")
    return out


def write_excel(summary: pd.DataFrame, today: date, files: dict) -> Path:
    plist = periods(today)
    table = pd.concat([overall(summary, columns(plist)), summary], ignore_index=True)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"{SUMMARY_NAME}_{today - timedelta(days=1):%d-%m-%Y}.xlsx"
    try:
        _write(table, path, plist, files)
    except PermissionError:
        path = path.with_name(f"{path.stem}_{int(time.time())}.xlsx")
        log(f"[WARN] Target is open in Excel - saving as {path.name}")
        _write(table, path, plist, files)
    return path


def _write(table: pd.DataFrame, path: Path, plist: list[Period], files: dict) -> None:
    labels = columns(plist)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        table.to_excel(xw, sheet_name="Master Summary", index=False, startrow=3)
        pd.DataFrame(
            [{"Report": n, "Period": p, "File": f.name} for (n, p), f in files.items()]
        ).to_excel(xw, sheet_name="Sources", index=False)

    wb = load_workbook(path)
    ws = wb["Master Summary"]
    ncols = 3 + len(labels)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    band_fill = PatternFill("solid", fgColor="DDEBF7")
    total_fill = PatternFill("solid", fgColor="FFF2CC")

    ws["A1"] = f"ASM-wise {SOURCE_LABEL} Enquiry > Test Ride > Booking > Retail"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = " | ".join(f"{p.label}: {p.start:%d-%b-%Y} to {p.end:%d-%b-%Y}" for p in plist)
    ws["A2"].font = Font(italic=True, color="595959")

    # Two header rows: the source groups (row 3) over the periods (row 4).
    for c in (1, 2, 3):
        ws.cell(row=3, column=c, value=ws.cell(row=4, column=c).value)
        ws.merge_cells(start_row=3, end_row=4, start_column=c, end_column=c)
    for i, (group, _) in enumerate(SOURCES):
        start = 4 + i * len(plist)
        ws.cell(row=3, column=start, value=group)
        ws.merge_cells(start_row=3, end_row=3, start_column=start, end_column=start + len(plist) - 1)
        for j, p in enumerate(plist):
            ws.cell(row=4, column=start + j, value=p.label)
    for r in (3, 4):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # First block is the All Dealers total; the rest are shaded alternately.
    first = 5
    n = len(METRICS)
    for block, start in enumerate(range(first, first + len(table), n)):
        end = start + n - 1
        fill = total_fill if block == 0 else (band_fill if block % 2 == 0 else None)
        for r in range(start, end + 1):
            metric = ws.cell(row=r, column=3).value
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                if fill:
                    cell.fill = fill
                if c > 3:
                    cell.number_format = "0.0" if metric in PERCENT_METRICS else "#,##0"
                    cell.alignment = Alignment(horizontal="right")
                    if metric in PERCENT_METRICS and cell.value is None:
                        cell.value = "-"
                        cell.alignment = Alignment(horizontal="center")
            if metric in PERCENT_METRICS:
                ws.cell(row=r, column=3).font = Font(italic=True)
        for c in (1, 2):
            ws.merge_cells(start_row=start, end_row=end, start_column=c, end_column=c)
            ws.cell(row=start, column=c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row=start, column=c).font = Font(bold=True)

    for c, width in zip(range(1, ncols + 1), [14, 26, 28, *[11] * len(labels)]):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "D5"
    ws.sheet_view.showGridLines = False
    wb["Sources"].column_dimensions["C"].width = 60
    wb.save(path)


def clear_downloads() -> None:
    """Empties the download folder, skipping anything still open."""
    root = DOWNLOAD_DIR
    if not root.is_dir():
        return
    removed = 0
    for item in root.iterdir():
        try:
            shutil.rmtree(item) if item.is_dir() and not item.is_symlink() else item.unlink()
            removed += 1
        except OSError as e:
            log(f"[WARN] Could not delete {item}: {e}")
    log(f"Cleared {removed} item(s) from {root}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--skip-download", action="store_true", help="rebuild from the last downloaded exports")
    ap.add_argument("--keep-downloads", action="store_true", help="keep the exports under Downloads")
    args = ap.parse_args()
    today = date.fromisoformat(os.environ["ASM_TODAY"]) if os.environ.get("ASM_TODAY") else date.today()
    for p in periods(today):
        log(f"{p.label}: {p.start} .. {p.end}")

    files = latest_files(today) if args.skip_download else download_reports(today)
    summary = build_summary(files, today)
    out = write_excel(summary, today, files)
    log(f"Master Summary: {out}  ({summary[['Dealer Code', 'ASM']].drop_duplicates().shape[0]} dealer/ASM pairs)")
    if not args.keep_downloads:
        clear_downloads()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
